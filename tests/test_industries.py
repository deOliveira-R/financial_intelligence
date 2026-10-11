import io
import zipfile
from types import SimpleNamespace

import openpyxl
import pytest
import respx
from fastapi.testclient import TestClient
from sqlalchemy import func, select

from fin_intel import graph, industries, themes
from fin_intel.api import app
from fin_intel.db import get_session, session_factory
from fin_intel.models import Concordance, GraphEdge, IoCode, IoNaics, IoUse, Issuer, Security
from fin_intel.providers import BeaProvider, CensusProvider
from fin_intel.providers.census import CONCORDANCES
from fin_intel.rebuild import rebuild

URL = "https://apps.bea.gov/industry/iTables%20Static%20Files/AllTablesSUP.zip"

# BEA's "NAICS Codes" sheet: sector in column A, summary B, underlying summary C, detail D
# (title E, NAICS G), after five heading rows.
HIERARCHY = [
    ("Industry Codes and Aggregations",),
    ("Estimates are available at four levels of detail",),
    (),
    ("BEA Industry Code", None, None, None, "Industry Title"),
    ("Sector", "Summary", "U.Summary", "Detail", None, "Notes", "Related 2017 NAICS Codes"),
    (23, "CONSTRUCTION"),
    (None, 23, "Construction"),
    (None, None, 23, "Construction"),
    (None, None, None, 233240, "Power and communication structures", None, "23*"),
    (31, "MANUFACTURING"),
    (None, 331, "Primary metals"),
    (None, None, 331, "Primary metals"),
    (None, None, None, 331410, "Nonferrous metal smelting", None, "331410, 33142-3"),
    (None, 333, "Machinery"),
    (None, None, None, 333242, "Semiconductor machinery manufacturing", None, 333242),
    (None, 334, "Computer and electronic products"),
    (None, None, None, 334413, "Semiconductor manufacturing", None, 334413),
    (None, "3364OT", "Other transportation equipment"),
    (None, None, None, 336411, "Aircraft manufacturing", None, 336411),
    (None, None, None, 336414, "Guided missile manufacturing", None, 336414),
    (54, "PROFESSIONAL"),
    (None, "5412OP", "Miscellaneous professional services"),
    (None, None, None, 541700, "Scientific research and development services", None, 5417),
    (None, None, None, "S00500", "Federal general government (defense)", None, "n.a."),
]
INDUSTRIES = ["233240", "331410", "333242", "334413", "336411", "541700"]
FINAL = ["F04000", "F06C00"]
DETAIL_FLOWS = {
    ("331410", "334413"): 10,
    ("331410", "336411"): 30,
    ("334413", "336411"): 20,
    ("334413", "334413"): 5,  # own purchases: inputs, but no edge
    ("541700", "336411"): 10,
    ("233240", "334413"): 5,
    ("S00300", "334413"): 5,  # noncomparable imports: an input, not a supplier
    ("334413", "F04000"): 40,
    ("336411", "F06C00"): 50,
    ("331410", "F04000"): 10,
}


def detail_sheet(book):
    ws = book.create_sheet("2017")
    columns = INDUSTRIES + ["T001"] + FINAL + ["T019"]
    for row in (("The Use Table",), ("[Millions of dollars]",), ("BEA",), ()):
        ws.append(row)
    ws.append(["Commodities/Industries", None, *(f"title {c}" for c in columns)])
    ws.append(["Code", "Commodity Description", *(int(c) if c.isdigit() else c for c in columns)])
    for row in [*INDUSTRIES, "S00300", "T005", "V00100"]:
        cells = [DETAIL_FLOWS.get((row, c)) for c in columns]
        ws.append([int(row) if row.isdigit() else row, f"title {row}", *cells])
    ws.append(["Note.  Detail may not add to total due to rounding."])


def summary_sheet(book):
    ws = book.create_sheet("2023")
    columns = ["331", "334", "3364OT", "T001", "F040", "T019"]
    for row in (("The Use of Commodities by Industries - Summary",), (), (), ("2023",), ()):
        ws.append(row)
    ws.append([None, "Commodities/Industries", *columns])
    ws.append(["IOCode", "Name", *(f"name {c}" for c in columns)])
    ws.append(["331", "Primary metals", 1, 12, 30, 43, 10, 53])
    ws.append(["334", "Computer and electronic products", None, 5, "...", 5, 40, 45])
    ws.append(["3364OT", "Other transportation equipment", None, None, 4, 4, 50, 54])


def bea_zip():
    detail = openpyxl.Workbook()
    detail.remove(detail.active)
    naics = detail.create_sheet("NAICS Codes")
    for row in HIERARCHY:
        naics.append(row)
    detail_sheet(detail)
    summary = openpyxl.Workbook()
    summary.remove(summary.active)
    summary_sheet(summary)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        for name, book in ((industries.DETAIL_FILE, detail), (industries.SUMMARY_FILE, summary)):
            data = io.BytesIO()
            book.save(data)
            archive.writestr(name, data.getvalue())
    return out.getvalue()


# Census concordances (title rows, a header, then code pairs): missile making splits between
# manufacturing and R&D, and R&D splits again in later NAICS revisions.
STEPS = {
    "sic1987-naics2002": [
        (3674, "Semiconductors", 334413),
        (3761, "Guided missiles", 336414),
        (3761, "Guided missiles (research and development)", 541710),
        ("Aux", "Auxiliaries", 551114),
    ],
    "naics2002-naics2007": [(541710, "R&D", 541711), (541710, "R&D", 541712)],
    "naics2007-naics2012": [],
    "naics2012-naics2017": [(541711, "R&D", 541713), (541711, "R&D", 541714)],
}


def concordance(name):
    book = openpyxl.Workbook()
    ws = book.active
    ws.append([f"{name} (Full Concordance)"])
    ws.append(["Code", "Title", "Code", "Title"])
    for code, title, target in STEPS[name]:
        ws.append([code, title, target, "target title"])
    data = io.BytesIO()
    book.save(data)
    return data.getvalue()


def test_parse_naics():
    assert industries.parse_naics("1122, 1124-5, 1129") == ["1122", "1124", "1125", "1129"]
    assert industries.parse_naics("722514-5") == ["722514", "722515"]
    assert industries.parse_naics("5174-9") == ["5174", "5175", "5176", "5177", "5178", "5179"]
    assert industries.parse_naics("23*") == ["23"]
    assert industries.parse_naics(111200) == ["111200"]
    assert industries.parse_naics("n.a.") == []


@pytest.fixture
def loaded(session, raw_store, monkeypatch):
    monkeypatch.setattr(
        graph.universe, "read", lambda: [SimpleNamespace(cik=1, industry="semiconductors")]
    )
    session.add_all(
        [
            Issuer(cik=1, name="Chips", sic=3674),
            Issuer(cik=2, name="Missiles", sic=3760),  # SEC group code: 376x
            Issuer(cik=3, name="Blank check", sic=6770),
            Issuer(cik=4, name="Unlisted chips", sic=3674),
        ]
    )
    session.flush()
    session.add_all(
        [
            Security(ticker="CHIP", cik=1),
            Security(ticker="MSL", cik=2),
            Security(ticker="SPAC", cik=3),
        ]
    )
    session.commit()
    with respx.mock:
        route = respx.get(URL).respond(content=bea_zip())
        for name, path in CONCORDANCES.items():
            respx.get(f"https://www.census.gov{path}").respond(content=concordance(name))
        bea, census = BeaProvider(raw_store=raw_store), CensusProvider(raw_store=raw_store)
        industries.sync_industries(session, bea, census)
        industries.sync_industries(session, bea, census)  # synced this month: no download
        assert route.call_count == 1
    return session


def test_load_tables(loaded):
    session = loaded
    codes = {(c.level, c.code): c for c in session.scalars(select(IoCode))}
    assert codes[("detail", "334413")].summary == "334"
    assert codes[("detail", "334413")].sector == "31"
    assert codes[("detail", "F06C00")].kind == "final_demand"
    assert codes[("detail", "S00300")].kind == "commodity"
    assert codes[("detail", "V00100")].kind == "value_added"
    assert codes[("summary", "3364OT")].kind == "industry"
    naics = set(session.execute(select(IoNaics.code, IoNaics.naics)).all())
    assert {("331410", "33142"), ("331410", "33143"), ("233240", "23")} <= naics
    assert not any(code == "S00500" for code, _ in naics)
    cell = session.get(IoUse, ("detail", 2017, "331410", "336411"))
    assert cell.value == 30
    # Suppressed ("...") and empty cells are left out.
    assert session.get(IoUse, ("summary", 2023, "334", "3364OT")) is None
    assert session.get(IoUse, ("summary", 2023, "331", "3364OT")).value == 30


def test_graph_edges_and_shares(loaded):
    session = loaded
    edge = session.get(GraphEdge, ("industry:331410", "industry:336411", "supplies"))
    # Aircraft inputs: 30 metals + 20 chips + 10 R&D; metals output: 30 + 10 + 10 exports.
    assert edge.buyer_share == pytest.approx(0.5) and edge.supplier_share == pytest.approx(0.6)
    chips = session.get(GraphEdge, ("industry:331410", "industry:334413", "supplies"))
    # Chip inputs include own purchases and imports: 10 + 5 + 5 + 5.
    assert chips.buyer_share == pytest.approx(0.4)
    assert session.get(GraphEdge, ("industry:334413", "industry:334413", "supplies")) is None
    defense = session.get(GraphEdge, ("industry:336411", "demand:F06C00", "supplies"))
    assert defense.buyer_share is None and defense.supplier_share == pytest.approx(1.0)
    assert session.get(GraphEdge, ("industry:334413", "summary:334", "part_of"))
    assert session.get(GraphEdge, ("summary:331", "summary:3364OT", "supplies")).value == 30

    up = {r["node"]: r for r in graph.walk(session, "industry:336411", "upstream", 2, 0.1)}
    assert up["industry:331410"]["share"] == pytest.approx(0.5)  # best path: direct
    assert up["industry:233240"]["depth"] == 2
    assert up["industry:233240"]["share"] == pytest.approx(20 / 60 * 5 / 25)
    down = {r["node"] for r in graph.walk(session, "industry:331410", "downstream", 2, 0.1)}
    assert {"industry:336411", "demand:F06C00", "demand:F04000"} <= down
    with pytest.raises(ValueError):
        graph.walk(session, "industry:336411", "sideways")


def test_theme_layer(loaded):
    session = loaded
    members = graph.theme_industries(session)
    assert members["defense"] == {"336411": "full", "336414": "full", "541700": "partial"}
    assert members["semiconductors"] == {"334413": "full", "333242": "full"}
    assert members["ai_power"]["233240"] == "curated"
    assert "233240" not in members["shipping"]  # construction's "23" isn't a partial match
    note = session.get(GraphEdge, ("theme:defense", "industry:541700", "maps_to")).note
    assert note == "partial"

    edge = session.get(GraphEdge, ("theme:critical_minerals", "theme:defense", "supplies"))
    assert edge.value == 30 and edge.buyer_share == pytest.approx(0.5)
    assert session.get(GraphEdge, ("theme:semiconductors", "theme:defense", "supplies")).value == 20
    # The partially mapped R&D industry stays out of defense's aggregates.
    assert (
        session.scalar(
            select(func.count()).where(
                GraphEdge.dst == "theme:defense", GraphEdge.src.like("%541%")
            )
        )
        == 0
    )
    assert graph.resolve(session, "defense") == "theme:defense"
    assert graph.resolve(session, "334413") == "industry:334413"


def test_company_layer(loaded):
    session = loaded
    pairs = set(session.execute(select(Concordance.code, Concordance.target_code)).all())
    assert ("3674", "334413") in pairs and not any(code == "Aux" for code, _ in pairs)

    def maps(cik):
        rows = session.execute(
            select(GraphEdge.dst, GraphEdge.value).where(
                GraphEdge.src == f"company:{cik}", GraphEdge.kind == "maps_to"
            )
        ).all()
        return {dst: pytest.approx(value) for dst, value in rows}

    assert maps(1) == {"industry:334413": 1.0}
    # 3760 isn't a SIC code: its group's 3761 splits half manufacturing, half R&D (R&D's
    # later splits don't add weight).
    assert maps(2) == {"industry:336414": 0.5, "industry:541700": 0.5}
    assert (
        maps(3) == {}
        and session.get(GraphEdge, ("company:3", "theme:semiconductors", "part_of")) is None
    )
    assert graph.resolve(session, "company:1")
    with pytest.raises(ValueError):
        graph.resolve(session, "company:4")  # no listed security
    theme = graph.companies(session, "theme:semiconductors")
    assert [(c["node"], c["note"]) for c in theme] == [("company:1", "universe")]
    missiles = graph.companies(session, "industry:336414")
    assert missiles[0]["label"] == "Missiles (MSL)" and missiles[0]["weight"] == 0.5


def test_themes_for():
    assert themes.themes_for(1311, "energy_materials") == ["oil_gas"]
    assert themes.themes_for(1040, "energy_materials") == ["critical_minerals"]
    assert themes.themes_for(3674, "semiconductors") == ["semiconductors"]
    assert themes.themes_for(7372) == ["technology"]
    assert themes.themes_for(None) == []


def test_rebuild_reproduces(loaded, raw_store):
    session = loaded

    def snapshot():
        return session.execute(
            select(GraphEdge.src, GraphEdge.dst, GraphEdge.kind, GraphEdge.value).order_by(
                GraphEdge.src, GraphEdge.dst, GraphEdge.kind
            )
        ).all()

    before = snapshot()
    assert before
    with respx.mock:
        rebuild(session, raw_store, "industries")
    assert snapshot() == before


def test_api_graph_and_themes(loaded, engine):
    def override():
        with session_factory(engine)() as s:
            yield s

    app.dependency_overrides[get_session] = override
    client = TestClient(app)
    assert any(t["name"] == "defense" for t in client.get("/themes").json())
    theme = client.get("/themes/defense").json()
    assert theme["naics"][0] == "3364" and theme["graph"]["supplies_in"]
    assert client.get("/themes/nope").status_code == 404
    up = client.get("/graph/336411", params={"direction": "upstream", "min_share": 0.1}).json()
    assert up["node"] == "industry:336411"
    assert up["upstream"][0]["node"] == "industry:331410"
    assert "maps_to_out" in client.get("/graph/defense").json()["edges"]
    assert client.get("/graph/999999").status_code == 404
    found = client.get("/graph/semiconductors", params={"direction": "companies"}).json()
    assert found["companies"][0]["label"] == "Chips (CHIP)"
    app.dependency_overrides.clear()
